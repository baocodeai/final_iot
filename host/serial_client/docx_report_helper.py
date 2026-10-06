"""Script to generate the complete, publication-grade academic DOCX report.

Project: Triển khai và đánh giá bộ mã hóa TinyCAE lượng tử hóa INT8 cho nén ảnh trên ESP32
Author: Nguyễn Vũ Bảo (MSSV: 23110079)
Instructor: Thầy Hồ Nhựt Minh
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, List, Dict

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path("D:/projects/iot-final")
RESULTS_DIR = PROJECT_ROOT / "results"
FIGURES_DIR = RESULTS_DIR / "figures"
EVIDENCE_DIR = PROJECT_ROOT / "submission_package/07_hardware_evidence"
REPORT_OUTPUT_PATH = PROJECT_ROOT / "docs/reports/BAO_CAO_HOAN_CHINH_ESP32_TINYCAE.docx"
SUBMISSION_COPY_PATH = PROJECT_ROOT / "submission_package/01_report_slide/BAO_CAO_HOAN_CHINH_ESP32_TINYCAE.docx"

# Color constants
COLOR_NAVY = RGBColor(0x1F, 0x4E, 0x79)       # #1F4E79 Primary Headings
COLOR_SLATE = RGBColor(0x2E, 0x40, 0x53)      # #2E4053 Secondary Headings
COLOR_STEEL = RGBColor(0x2C, 0x3E, 0x50)      # #2C3E50 Tertiary Headings
COLOR_CHARCOAL = RGBColor(0x22, 0x22, 0x22)   # #222222 Body Text
COLOR_MUTED = RGBColor(0x55, 0x55, 0x55)      # #555555 Captions, footers
HEX_HEADER_BG = "1F4E79"
HEX_ZEBRA_BG = "F4F6F9"
HEX_CALLOUT_BG = "EBF5FB"
HEX_BORDER = "D0D7DE"

def read_csv(path: Path) -> List[Dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))

def set_cell_shading(cell, hex_color: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_table(table, col_widths: List[float] = None, alignment=WD_TABLE_ALIGNMENT.CENTER):
    table.alignment = alignment
    set_table_borders(table)
    
    # Format header row
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for cell in table.rows[0].cells:
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    # Format data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        # Zebra striping
        bg_hex = HEX_ZEBRA_BG if (r_idx % 2 == 0) else "FFFFFF"
        for cell in row.cells:
            if bg_hex != "FFFFFF":
                set_cell_shading(cell, bg_hex)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = COLOR_CHARCOAL
                    
    # Set widths if specified
    if col_widths:
        for row in table.rows:
            for c_idx, width in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = Inches(width)

def add_callout(doc, title: str, text: str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_CALLOUT_BG}"/>')
    tcPr.append(shd)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{HEX_HEADER_BG}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title.upper()}\n")
    run_t.font.name = "Times New Roman"
    run_t.font.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = COLOR_NAVY
    
    run_body = p.add_run(text)
    run_body.font.name = "Times New Roman"
    run_body.font.size = Pt(10)
    run_body.font.color.rgb = COLOR_CHARCOAL
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_1(doc, text: str):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return h

def add_heading_2(doc, text: str):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE
    return h

def add_heading_3(doc, text: str):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = COLOR_STEEL
    return h

def add_body(doc, text: str, bold_prefix: str = None, space_after: float = 4.0):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_CHARCOAL
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_CHARCOAL
    return p

def add_bullet(doc, text: str, bold_prefix: str = None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_CHARCOAL
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_CHARCOAL
    return p

def add_caption(doc, text: str, is_table: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_before = Pt(4) if not is_table else Pt(6)
    p.paragraph_format.space_after = Pt(6) if not is_table else Pt(3)
    p.paragraph_format.keep_with_next = True if is_table else False
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = COLOR_MUTED
    return p

def add_figure(doc, img_path: Path, caption: str, width_in: float = 6.0):
    if img_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        r = p_img.add_run()
        r.add_picture(str(img_path), width=Inches(width_in))
        add_caption(doc, caption, is_table=False)
    else:
        print(f"Warning: Figure not found: {img_path}")

print("Docx builder module initialized.")
